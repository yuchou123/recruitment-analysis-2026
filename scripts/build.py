# -*- coding: utf-8 -*-
"""生成《招聘项目三表对比分析.xlsx》：data/game/operator 三表分析汇总报表。

用法：
    python build.py [输出xlsx路径]
数据默认取仓库 data/analysis_data.json（可用环境变量 DATA_JSON 覆盖）。
"""
import json
import math
import os
import sys

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(BASE)
DATA_JSON = os.environ.get("DATA_JSON", os.path.join(REPO_ROOT, "data", "analysis_data.json"))
OUT_XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO_ROOT, "analysis", "招聘项目三表对比分析.xlsx")
DATA = json.load(open(DATA_JSON, encoding="utf-8"))

# ---------------- 颜色 / 样式常量 ----------------
def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value

XL_PRIMARY = xl_color("#4472C4")   # 表头蓝
XL_DARK = xl_color("#2F5597")      # 标题深蓝
XL_LIGHT = xl_color("#D9E2F3")     # 浅蓝底
XL_BORDER = xl_color("#BFBFBF")
XL_GRAY = xl_color("#808080")
XL_TEXT = xl_color("#1A1A1A")

FONT = "微软雅黑"
F_TITLE = Font(name=FONT, size=15, bold=True, color="FFFFFF")
F_SUB = Font(name=FONT, size=9, color=XL_GRAY)
F_SECTION = Font(name=FONT, size=11, bold=True, color=XL_DARK)
F_HEADER = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10, color=XL_TEXT)
F_BOLD = Font(name=FONT, size=10, bold=True, color=XL_TEXT)
F_KPI_L = Font(name=FONT, size=9, color=XL_GRAY)
F_KPI_V = Font(name=FONT, size=18, bold=True, color=XL_DARK)
F_NOTE = Font(name=FONT, size=8, italic=True, color=XL_GRAY)

FILL_TITLE = PatternFill("solid", fgColor=XL_DARK)
FILL_HEADER = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_LIGHT = PatternFill("solid", fgColor=XL_LIGHT)
FILL_TOTAL = PatternFill("solid", fgColor=XL_LIGHT)

thin = Side(style="thin", color=XL_BORDER)
thin_blue = Side(style="thin", color=XL_PRIMARY)
B_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
B_CARD = Border(left=thin_blue, right=thin_blue, top=thin_blue, bottom=thin_blue)

A_C = Alignment(horizontal="center", vertical="center")
A_L = Alignment(horizontal="left", vertical="center")
A_R = Alignment(horizontal="right", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

FMT_INT = "#,##0"
FMT_PCT1 = "0.0%"
FMT_PCT2 = "0.00%"

JOBS = [("data", "数据岗"), ("game", "游戏岗"), ("operator", "运维岗")]
SERIES_COLORS = ["4472C4", "ED7D31", "70AD47"]  # 图表三岗配色（不带 #）


def box(ws, cell_range, border):
    """对合并/普通区域的每个单元格补全边框。"""
    for row in ws[cell_range]:
        for c in row:
            c.border = border


def title_band(ws, last_col, text):
    """仅第 1 行标题带；第 2 行留给表头（调用方自行决定是否加副标题行）。"""
    lc = get_column_letter(last_col)
    ws.merge_cells(f"A1:{lc}1")
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A1"].fill = FILL_TITLE
    ws["A1"].alignment = A_C
    ws.row_dimensions[1].height = 30
    box(ws, f"A1:{lc}1", B_ALL)


def subtitle_band(ws, last_col, text):
    """第 2 行副标题（调用方需确保第 2 行无其他内容）。"""
    lc = get_column_letter(last_col)
    ws.merge_cells(f"A2:{lc}2")
    ws["A2"] = text
    ws["A2"].font = F_SUB
    ws["A2"].fill = FILL_LIGHT
    ws["A2"].alignment = A_C
    ws.row_dimensions[2].height = 18
    box(ws, f"A2:{lc}2", B_ALL)


def note_row(ws, row, last_col, text):
    lc = get_column_letter(last_col)
    ws.merge_cells(f"A{row}:{lc}{row}")
    ws[f"A{row}"] = text
    ws[f"A{row}"].font = F_NOTE
    ws[f"A{row}"].alignment = A_L
    ws.row_dimensions[row].height = 15


def header_row(ws, row, headers, start_col=1, height=20):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = F_HEADER
        c.fill = FILL_HEADER
        c.alignment = A_C
        c.border = B_ALL
    ws.row_dimensions[row].height = height


def section_label(ws, row, last_col, text):
    lc = get_column_letter(last_col)
    ws.merge_cells(f"A{row}:{lc}{row}")
    ws[f"A{row}"] = text
    ws[f"A{row}"].font = F_SECTION
    ws[f"A{row}"].fill = FILL_LIGHT
    ws[f"A{row}"].alignment = A_L
    ws.row_dimensions[row].height = 20
    box(ws, f"A{row}:{lc}{row}", B_ALL)


def put(ws, coord, value, font=F_BODY, align=A_R, fmt=None, fill=None, border=B_ALL):
    c = ws[coord]
    c.value = value
    c.font = font
    c.alignment = align
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def add_databars(ws, ranges, color="4472C4"):
    for rng in ranges:
        ws.conditional_formatting.add(
            rng,
            DataBarRule(start_type="num", start_value=0, end_type="max",
                        color=color, showValue=True),
        )


def color_series(chart, line=False):
    for s, color in zip(chart.series, SERIES_COLORS):
        if line:
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 22000
            s.marker = Marker(symbol="circle", size=6)
            s.smooth = False
        else:
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.noFill = True


wb = openpyxl.Workbook()
wb.properties.title = "招聘项目三表对比分析"

# ============================================================
# Sheet 2 先建（被其他表公式引用）：城市需求对比
# 锚点：R1 标题 / R2 表头 / R3-6 数据 / R7 合计
# ============================================================
ws_city = wb.active
ws_city.title = "城市需求对比"
title_band(ws_city, 7, "北上广深城市需求对比（职位数）")
header_row(ws_city, 2, ["城市", "数据岗职位数", "数据岗占比", "游戏岗职位数",
                        "游戏岗占比", "运维岗职位数", "运维岗占比"])

city_totals = {j: sum(DATA["city"][j].values()) for j, _ in JOBS}
city_order = sorted(DATA["city"]["data"].keys(),
                    key=lambda c: -(DATA["city"]["data"][c] + DATA["city"]["game"][c] + DATA["city"]["operator"][c]))
for i, city in enumerate(city_order):
    r = 3 + i
    put(ws_city, f"A{r}", city, font=F_BODY, align=A_C)
    for col, job in (("B", "data"), ("D", "game"), ("F", "operator")):
        put(ws_city, f"{col}{r}", DATA["city"][job][city], fmt=FMT_INT)
    put(ws_city, f"C{r}", f'=IF(OR(B$7="",B$7=0),"",B{r}/B$7)', fmt=FMT_PCT2)
    put(ws_city, f"E{r}", f'=IF(OR(D$7="",D$7=0),"",D{r}/D$7)', fmt=FMT_PCT2)
    put(ws_city, f"G{r}", f'=IF(OR(F$7="",F$7=0),"",F{r}/F$7)', fmt=FMT_PCT2)

put(ws_city, "A7", "合计", font=F_BOLD, align=A_C, fill=FILL_TOTAL)
put(ws_city, "B7", "=SUM(B3:B6)", font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)
put(ws_city, "C7", '=IF(OR(B7="",B7=0),"",B7/B7)', font=F_BOLD, fmt=FMT_PCT2, fill=FILL_TOTAL)
put(ws_city, "D7", "=SUM(D3:D6)", font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)
put(ws_city, "E7", '=IF(OR(D7="",D7=0),"",D7/D7)', font=F_BOLD, fmt=FMT_PCT2, fill=FILL_TOTAL)
put(ws_city, "F7", "=SUM(F3:F6)", font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)
put(ws_city, "G7", '=IF(OR(F7="",F7=0),"",F7/F7)', font=F_BOLD, fmt=FMT_PCT2, fill=FILL_TOTAL)

for col, w in zip("ABCDEFG", [10, 13, 11, 13, 11, 13, 11]):
    ws_city.column_dimensions[col].width = w
add_databars(ws_city, ["B3:B6", "D3:D6", "F3:F6"])
note_row(ws_city, 8, 7, "口径：v_*_clean_workplace 视图 · 按城市汇总的最终分析集职位数；占比列为公式计算。")

ch = BarChart()
ch.type = "col"
ch.grouping = "clustered"
ch.title = "北上广深职位数对比"
ch.y_axis.title = "职位数"
ch.x_axis.title = "城市"
ch.height = 8.5
ch.width = 16
for col in (2, 4, 6):
    ch.add_data(Reference(ws_city, min_col=col, min_row=2, max_row=6), titles_from_data=True)
ch.set_categories(Reference(ws_city, min_col=1, min_row=3, max_row=6))
ch.legend.position = "b"
color_series(ch)
ws_city.add_chart(ch, "A9")

# ============================================================
# Sheet 3：企业类型分析
# 锚点：R1 标题 / R2 表头 / R3-12 数据 / R13 合计
# ============================================================
ws_ctype = wb.create_sheet("企业类型分析")
title_band(ws_ctype, 10, "企业类型分析（招聘量与平均月薪）")
header_row(ws_ctype, 2, ["企业类型", "数据岗招聘量", "数据岗占比", "游戏岗招聘量", "游戏岗占比",
                         "运维岗招聘量", "运维岗占比", "数据岗平均月薪（元）",
                         "游戏岗平均月薪（元）", "运维岗平均月薪（元）"], height=32)

ctype_map = {j: {r["type"]: r for r in DATA["ctype"][j]} for j, _ in JOBS}
all_types = sorted(
    {t for j, _ in JOBS for t in ctype_map[j]},
    key=lambda t: -sum(ctype_map[j].get(t, {}).get("count", 0) for j, _ in JOBS),
)
for i, t in enumerate(all_types):
    r = 3 + i
    put(ws_ctype, f"A{r}", t, font=F_BODY, align=A_C)
    for col, job in (("B", "data"), ("D", "game"), ("F", "operator")):
        info = ctype_map[job].get(t)
        put(ws_ctype, f"{col}{r}", info["count"] if info else None, fmt=FMT_INT)
    put(ws_ctype, f"C{r}", f'=IF(OR(B$13="",B$13=0),"",B{r}/B$13)', fmt=FMT_PCT2)
    put(ws_ctype, f"E{r}", f'=IF(OR(D$13="",D$13=0),"",D{r}/D$13)', fmt=FMT_PCT2)
    put(ws_ctype, f"G{r}", f'=IF(OR(F$13="",F$13=0),"",F{r}/F$13)', fmt=FMT_PCT2)
    for col, job in (("H", "data"), ("I", "game"), ("J", "operator")):
        info = ctype_map[job].get(t)
        put(ws_ctype, f"{col}{r}", info["avg_salary"] if info and info["avg_salary"] is not None else None,
            fmt=FMT_INT)

put(ws_ctype, "A13", "合计 / 整体", font=F_BOLD, align=A_C, fill=FILL_TOTAL)
for col in ("B", "D", "F"):
    put(ws_ctype, f"{col}13", f"=SUM({col}3:{col}12)", font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)
for col, base, sal in (("C", "B", "B"), ("E", "D", "D"), ("G", "F", "F")):
    put(ws_ctype, f"{col}13", f'=IF(OR({base}13="",{base}13=0),"",{base}13/{base}13)',
        font=F_BOLD, fmt=FMT_PCT2, fill=FILL_TOTAL)
for col, base, sal in (("H", "B", "H"), ("I", "D", "I"), ("J", "F", "J")):
    put(ws_ctype, f"{col}13",
        f'=IF(OR(SUM({base}3:{base}12)=0,COUNT({sal}3:{sal}12)=0),"",'
        f'ROUND(SUMPRODUCT({base}3:{base}12,{sal}3:{sal}12)/SUM({base}3:{base}12),0))',
        font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)

for col, w in zip("ABCDEFGHIJ", [14, 12, 10, 12, 10, 12, 10, 17, 17, 17]):
    ws_ctype.column_dimensions[col].width = w
add_databars(ws_ctype, ["B3:B12", "D3:D12", "F3:F12"])
note_row(ws_ctype, 14, 10, "口径：招聘量 = COUNT(*) · 平均月薪 = 该类型 avg(salary_mean) · 占比与合计均为公式计算。")

ch2 = BarChart()
ch2.type = "col"
ch2.grouping = "clustered"
ch2.title = "各企业类型平均月薪对比"
ch2.y_axis.title = "平均月薪（元）"
ch2.height = 9
ch2.width = 20
for col in (8, 9, 10):
    ch2.add_data(Reference(ws_ctype, min_col=col, min_row=2, max_row=12), titles_from_data=True)
ch2.set_categories(Reference(ws_ctype, min_col=1, min_row=3, max_row=12))
ch2.legend.position = "b"
color_series(ch2)
ws_ctype.add_chart(ch2, "A15")

# ============================================================
# Sheet 4：年限薪资对比
# 锚点：R1 标题 / R2 表头 / R3-10 数据 / R11 整体平均 / R12 注 / 图表 A14
# ============================================================
ws_wy = wb.create_sheet("年限薪资对比")
title_band(ws_wy, 7, "工作年限 - 平均月薪对比")
header_row(ws_wy, 2, ["工作年限（年）", "数据岗平均月薪（元）", "游戏岗平均月薪（元）",
                      "运维岗平均月薪（元）", "数据岗样本数", "游戏岗样本数", "运维岗样本数"], height=32)

years = sorted({y for j, _ in JOBS for y in DATA["workyear"][j]}, key=lambda s: (len(s), s))
for i, y in enumerate(years):
    r = 3 + i
    put(ws_wy, f"A{r}", int(y), font=F_BODY, align=A_C, fmt="0")
    for col, job in (("B", "data"), ("C", "game"), ("D", "operator")):
        put(ws_wy, f"{col}{r}", DATA["workyear"][job][y], fmt=FMT_INT)
    for col, job in (("E", "data"), ("F", "game"), ("G", "operator")):
        put(ws_wy, f"{col}{r}", DATA["workyear_count"][job].get(y), fmt=FMT_INT)

put(ws_wy, "A11", "整体平均", font=F_BOLD, align=A_C, fill=FILL_TOTAL)
for col, sal, cnt in (("B", "B", "E"), ("C", "C", "F"), ("D", "D", "G")):
    put(ws_wy, f"{col}11",
        f'=IF(OR(SUM({cnt}3:{cnt}10)=0,COUNT({sal}3:{sal}10)=0),"",'
        f'ROUND(SUMPRODUCT({sal}3:{sal}10,{cnt}3:{cnt}10)/SUM({cnt}3:{cnt}10),0))',
        font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)
for col in ("E", "F", "G"):
    put(ws_wy, f"{col}11", f"=SUM({col}3:{col}10)", font=F_BOLD, fmt=FMT_INT, fill=FILL_TOTAL)

ws_wy.merge_cells("A12:G12")
put(ws_wy, "A12", "注：样本数为该年限下纳入统计的职位数；样本 < 10 时均值波动极大（如数据岗「8 年」仅 1 条），仅供参考。",
    font=F_NOTE, align=A_L, border=None)

for col, w in zip("ABCDEFG", [14, 17, 17, 17, 12, 12, 12]):
    ws_wy.column_dimensions[col].width = w

ch3 = LineChart()
ch3.title = "工作年限 - 平均月薪走势"
ch3.y_axis.title = "平均月薪（元）"
ch3.x_axis.title = "工作年限（年）"
ch3.height = 8.5
ch3.width = 16
for col in (2, 3, 4):
    ch3.add_data(Reference(ws_wy, min_col=col, min_row=2, max_row=10), titles_from_data=True)
ch3.set_categories(Reference(ws_wy, min_col=1, min_row=3, max_row=10))
ch3.legend.position = "b"
color_series(ch3, line=True)
ws_wy.add_chart(ch3, "A14")

# ============================================================
# Sheet 5：核心技能对比
# 锚点：R1 标题 / R2 表头 / R3-17 数据（Top15） / 图表 A19
# ============================================================
ws_skill = wb.create_sheet("核心技能对比")
title_band(ws_skill, 7, "核心技能出现频率对比（Top 15）")
header_row(ws_skill, 2, ["技能点", "数据岗频数", "游戏岗频数", "运维岗频数",
                         "数据岗出现频率", "游戏岗出现频率", "运维岗出现频率"])

skills = sorted(DATA["skills_all"],
                key=lambda s: -max(s["data"], s["game"], s["operator"]))[:15]
for i, s in enumerate(skills):
    r = 3 + i
    put(ws_skill, f"A{r}", s["skill"], font=F_BODY, align=A_C)
    put(ws_skill, f"B{r}", s["data"], fmt=FMT_INT)
    put(ws_skill, f"C{r}", s["game"], fmt=FMT_INT)
    put(ws_skill, f"D{r}", s["operator"], fmt=FMT_INT)
    put(ws_skill, f"E{r}", "=IF(OR('城市需求对比'!$B$7=\"\",'城市需求对比'!$B$7=0),\"\","
                           f"B{r}/'城市需求对比'!$B$7)", fmt=FMT_PCT2)
    put(ws_skill, f"F{r}", "=IF(OR('城市需求对比'!$D$7=\"\",'城市需求对比'!$D$7=0),\"\","
                           f"C{r}/'城市需求对比'!$D$7)", fmt=FMT_PCT2)
    put(ws_skill, f"G{r}", "=IF(OR('城市需求对比'!$F$7=\"\",'城市需求对比'!$F$7=0),\"\","
                           f"D{r}/'城市需求对比'!$F$7)", fmt=FMT_PCT2)

for col, w in zip("ABCDEFG", [12, 11, 11, 11, 14, 14, 14]):
    ws_skill.column_dimensions[col].width = w
add_databars(ws_skill, ["E3:E17", "F3:F17", "G3:G17"])
note_row(ws_skill, 18, 7, "口径：频数 = 职位详情 LIKE 匹配数（全量技能，已剔除「报告」） · 频率 = 频数 / 该岗职位数（公式）。")

ch4 = BarChart()
ch4.type = "bar"
ch4.grouping = "clustered"
ch4.title = "核心技能出现频率对比"
ch4.height = 14
ch4.width = 18
for col in (5, 6, 7):
    ch4.add_data(Reference(ws_skill, min_col=col, min_row=2, max_row=17), titles_from_data=True)
ch4.set_categories(Reference(ws_skill, min_col=1, min_row=3, max_row=17))
ch4.legend.position = "b"
color_series(ch4)
ws_skill.add_chart(ch4, "A19")

# ============================================================
# Sheet 1：总览看板（引用以上各表，公式联动）
# ============================================================
ws = wb.create_sheet("总览看板", 0)
title_band(ws, 8, "2026年招聘市场三岗位对比分析")
subtitle_band(ws, 8, "数据来源：recruitment_2026 数据库 data / game / operator 三表 · 统计口径详见「说明」表 · 生成日期：2026-10-26")

section_label(ws, 4, 8, "一、核心指标")
cards = [
    ("A5:B5", "A6:B6", "数据岗职位数", "=SUM('城市需求对比'!B3:B6)"),
    ("C5:D5", "C6:D6", "游戏岗职位数", "=SUM('城市需求对比'!D3:D6)"),
    ("E5:F5", "E6:F6", "运维岗职位数", "=SUM('城市需求对比'!F3:F6)"),
    ("G5:H5", "G6:H6", "三岗职位合计", "=A6+C6+E6"),
]
for lab_rng, val_rng, label, formula in cards:
    ws.merge_cells(lab_rng)
    ws.merge_cells(val_rng)
    lc = ws[lab_rng.split(":")[0]]
    lc.value = label
    lc.font = F_KPI_L
    lc.alignment = A_C
    vc = ws[val_rng.split(":")[0]]
    vc.value = formula
    vc.font = F_KPI_V
    vc.alignment = A_C
    vc.number_format = FMT_INT
    for row in ws[val_rng]:
        for c in row:
            c.fill = FILL_LIGHT
    box(ws, f"{lab_rng.split(':')[0]}:{val_rng.split(':')[1]}", B_CARD)
ws.row_dimensions[5].height = 16
ws.row_dimensions[6].height = 32

section_label(ws, 8, 8, "二、清洗漏斗对比（各阶段记录数）")
header_row(ws, 9, ["阶段", "数据岗", "游戏岗", "运维岗",
                   "数据岗保留率", "游戏岗保留率", "运维岗保留率", "阶段说明"])
funnel_rows = [
    ("原始记录", "data_raw", "game_raw", "op_raw", "三表全量数据"),
    ("清洗空值后", "data_null", "game_null", "op_null", "13 个关键字段非空"),
    ("去重后", "data_dist", "game_dist", "op_dist", "公司 + 职位取最新"),
    ("限北上广深", "data_wp", "game_wp", "op_wp", "排除其他城市"),
    ("岗位筛选后", "data_final", "game_final", "op_final", "%数据% / %游戏% / %运维%"),
]
for i, (stage, kd, kg, ko, note) in enumerate(funnel_rows):
    r = 10 + i
    bold = (i == len(funnel_rows) - 1)
    f = F_BOLD if bold else F_BODY
    fill = FILL_TOTAL if bold else None
    put(ws, f"A{r}", stage, font=f, align=A_C, fill=fill)
    put(ws, f"B{r}", DATA["funnel"][kd], font=f, fmt=FMT_INT, fill=fill)
    put(ws, f"C{r}", DATA["funnel"][kg], font=f, fmt=FMT_INT, fill=fill)
    put(ws, f"D{r}", DATA["funnel"][ko], font=f, fmt=FMT_INT, fill=fill)
    put(ws, f"E{r}", f'=IF(OR(B$10="",B$10=0),"",B{r}/B$10)', font=f, fmt=FMT_PCT1, fill=fill)
    put(ws, f"F{r}", f'=IF(OR(C$10="",C$10=0),"",C{r}/C$10)', font=f, fmt=FMT_PCT1, fill=fill)
    put(ws, f"G{r}", f'=IF(OR(D$10="",D$10=0),"",D{r}/D$10)', font=f, fmt=FMT_PCT1, fill=fill)
    put(ws, f"H{r}", note, font=F_NOTE, align=A_C, fill=fill)

section_label(ws, 16, 8, "三、三岗核心指标速览")
header_row(ws, 17, ["岗位", "职位数", "平均月薪（元）", "主力城市", "主导企业类型", "TOP技能"])
for i, (job, name) in enumerate(JOBS):
    r = 18 + i
    # cb=城市需求对比表的职位数列(B/D/F)，cs=核心技能对比表的频数列(B/C/D)
    cb, cs = {"data": ("B", "B"), "game": ("D", "C"), "operator": ("F", "D")}[job]
    put(ws, f"A{r}", name, font=F_BOLD, align=A_C)
    put(ws, f"B{r}", f"=SUM('城市需求对比'!{cb}3:{cb}6)", fmt=FMT_INT)
    put(ws, f"C{r}",
        f'=IF(OR(SUM(\'企业类型分析\'!{cb}$3:{cb}$12)=0,COUNT(\'企业类型分析\'!{cb}$3:{cb}$12)=0),"",'
        f'ROUND(SUMPRODUCT(\'企业类型分析\'!{cb}$3:{cb}$12,\'企业类型分析\'!{"HIJ"[i]}$3:{"HIJ"[i]}$12)'
        f'/SUM(\'企业类型分析\'!{cb}$3:{cb}$12),0))', fmt=FMT_INT)
    put(ws, f"D{r}",
        f"=INDEX('城市需求对比'!$A$3:$A$6,MATCH(MAX('城市需求对比'!{cb}$3:{cb}$6),"
        f"'城市需求对比'!{cb}$3:{cb}$6,0))", align=A_C)
    put(ws, f"E{r}",
        f"=INDEX('企业类型分析'!$A$3:$A$12,MATCH(MAX('企业类型分析'!{cb}$3:{cb}$12),"
        f"'企业类型分析'!{cb}$3:{cb}$12,0))", align=A_C)
    put(ws, f"F{r}",
        f"=INDEX('核心技能对比'!$A$3:$A$17,MATCH(MAX('核心技能对比'!{cs}$3:{cs}$17),"
        f"'核心技能对比'!{cs}$3:{cs}$17,0))", align=A_C)

ws.merge_cells("A21:H21")
put(ws, "A21", "注：平均月薪按企业类型招聘量加权计算；主力城市 / 主导企业类型 / TOP技能 由公式取各维度最大值自动得出。",
    font=F_NOTE, align=A_L, border=None)

ch1 = BarChart()
ch1.type = "col"
ch1.title = "三岗平均月薪对比"
ch1.y_axis.title = "平均月薪（元）"
ch1.height = 7.5
ch1.width = 14
ch1.add_data(Reference(ws, min_col=3, min_row=17, max_row=20), titles_from_data=True)
ch1.set_categories(Reference(ws, min_col=1, min_row=18, max_row=20))
ch1.legend = None
color_series(ch1)
ws.add_chart(ch1, "A23")

for col, w in zip("ABCDEFGH", [16, 13, 15, 12, 14, 12, 13, 20]):
    ws.column_dimensions[col].width = w

# ============================================================
# Sheet 6：说明
# ============================================================
ws_doc = wb.create_sheet("说明")
title_band(ws_doc, 2, "说明")
header_row(ws_doc, 2, ["项目", "说明"])
notes = [
    ("数据来源", "本地 MySQL「recruitment_2026」数据库的 data（数据岗）、game（游戏岗）、operator（运维岗）三张表，原始记录 8,000 / 5,000 / 5,000 条，职位发布日期为 2026 年。"),
    ("岗位筛选口径", "与 v_*_clean 视图一致：数据岗 job_name LIKE '%数据%'、游戏岗 LIKE '%游戏%'、运维岗 LIKE '%运维%'。"),
    ("城市口径", "仅保留北上广深（workarea_text LIKE 匹配），其他城市与「异地招聘」已在清洗流程中排除。"),
    ("去重口径", "同一「公司名 + 职位名」仅保留 issuedate 最新的一条（ROW_NUMBER 窗口函数）。"),
    ("薪资口径", "providesalary_text 解析为月薪：千/月 ×1000、万/月 ×10000、万/年 ×833（约 ÷12）；salary_mean = (salary_min + salary_max) / 2；个别「1.5千以下/月」等非区间文本无法解析，计为 NULL（占比极小）。"),
    ("招聘量口径", "本报表「招聘量 / 职位数」均按 COUNT(*) 统计；SQL 练习视图中的「招聘总数」为 SUM(degreefrom)（学历编码求和），仅排序趋势一致，绝对值不代表职位数。"),
    ("技能口径", "对 v_*_clean 全量职位详情做 LIKE 匹配 skill_table 的 63 个技能关键词；「报告」一词因高频出现于「测试报告 / 运行报告」等文本、并非真实技能，本报表已剔除。"),
    ("口径统一调整", "v_data_salary_unit 原基于 v_clean_data_distinct（7,058 行），已按用户脚本改回基于 v_data_clean（495 行），与 game / operator 薪资链口径一致，其下游 3 个视图已原样重建。"),
    ("样本量提示", "年限薪资表中样本数 < 10 的格子（如数据岗「8 年」仅 1 条）波动极大，解读需谨慎。"),
    ("配套文件", "招聘项目查询.sql（data 表分析）、招聘项目查询1.sql（game）、招聘项目查询2.sql（operator）、game_operator_analysis.sql（26 个视图脚本）；本工作簿由 AI 于 2026-10-26 生成。"),
]
for i, (k, v) in enumerate(notes):
    r = 3 + i
    put(ws_doc, f"A{r}", k, font=F_BOLD, align=Alignment(horizontal="left", vertical="top"))
    put(ws_doc, f"B{r}", v, font=F_BODY, align=A_WRAP)
    ws_doc.row_dimensions[r].height = max(18, 15 * math.ceil(len(v) / 52) + 6)
ws_doc.column_dimensions["A"].width = 16
ws_doc.column_dimensions["B"].width = 100

# 标签颜色
for s, color in [("总览看板", "2F5597"), ("城市需求对比", "4472C4"), ("企业类型分析", "4472C4"),
                 ("年限薪资对比", "4472C4"), ("核心技能对比", "4472C4"), ("说明", "808080")]:
    wb[s].sheet_properties.tabColor = color

wb.save(OUT_XLSX)
print("saved:", OUT_XLSX)
print("sheets:", wb.sheetnames)
